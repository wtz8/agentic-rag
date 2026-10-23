"""Office 解析器：DOCX / XLSX / PPTX（03 §5）。

默认走 `markitdown` 转 Markdown（统一中间表示），不可用时降级到
`python-docx` / `openpyxl` / `python-pptx` 自行拼 Markdown——
降级必须被记录（见 `.via_fallback`），否则等于静默降质。
"""

from __future__ import annotations

import re
from pathlib import Path

from ...schema import ParsedDocument
from ..quality import markdown_to_blocks
from .base import ParserError, make_doc


def _markitdown_md(path: Path) -> str:
    try:
        from markitdown import MarkItDown
    except Exception as e:  # noqa: BLE001
        raise ParserError(f"markitdown 不可用: {e}") from e
    try:
        return MarkItDown().convert(str(path)).text_content or ""
    except Exception as e:  # noqa: BLE001
        raise ParserError(f"markitdown 解析失败: {e}") from e


# ── 降级实现（不依赖 markitdown）────────────────────────────────────────


def _docx_md(path: Path) -> str:
    import docx  # python-docx

    d = docx.Document(str(path))
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        m = re.search(r"heading\s*(\d)", style)
        if m:
            out.append("#" * min(6, int(m.group(1))) + " " + text)
        elif style.startswith("title"):
            out.append("# " + text)
        else:
            out.append(text)
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        out.append(_rows_to_md(rows))
    return "\n\n".join(out)


def _xlsx_md(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                rows.append(["" if v is None else str(v).strip() for v in row])
        if rows:
            out.append(_rows_to_md(rows))
    return "\n\n".join(out)


def _pptx_md(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"## 幻灯片 {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for line in shape.text_frame.text.splitlines():
                    if line.strip():
                        out.append(line.strip())
            if getattr(shape, "has_table", False):
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                out.append(_rows_to_md(rows))
            if shape.shape_type is not None and getattr(shape, "image", None) is not None:
                out.append(f"![](image_{i})")
    return "\n\n".join(out)


def _rows_to_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    lines = []
    for i, r in enumerate(rows):
        cells = list(r) + [""] * (width - len(r))
        lines.append("| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |")
        if i == 0:
            lines.append("| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines)


class _OfficeParser:
    name = "markitdown"
    _fallback = staticmethod(lambda path: "")

    def __init__(self) -> None:
        self.via_fallback = False

    def parse(self, path: Path) -> ParsedDocument:
        try:
            md = _markitdown_md(Path(path))
            self.via_fallback = False
        except Exception:  # noqa: BLE001 - 任何失败都降级到原生实现，但必须被记录
            md = self._fallback(Path(path))
            self.via_fallback = True
        if not md.strip():
            raise ParserError(f"{self.name} 解析结果为空: {path}")
        doc = make_doc(path, markdown_to_blocks(md), self.name)
        if self.via_fallback:
            doc.quality.issues.append("parser_fallback_native")
        return doc


class DocxParser(_OfficeParser):
    _fallback = staticmethod(_docx_md)


class XlsxParser(_OfficeParser):
    _fallback = staticmethod(_xlsx_md)


class PptxParser(_OfficeParser):
    _fallback = staticmethod(_pptx_md)


__all__ = ["DocxParser", "XlsxParser", "PptxParser"]
