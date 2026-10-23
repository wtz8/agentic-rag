"""解析质量三道关卡（03 §4）。

工程目标不是"解析 100% 无误"（做不到），而是**错误可发现、可定位、不扩散**。

| 关卡 | 检查项 | 失败处置 |
| --- | --- | --- |
| 解析前探测 | 扫描件、加密损坏、表格密度 | 扫描件 → OCR 通道 |
| 解析中结构校验 | 表格行列、标题层级、图文关联 | 换解析器 / 记 issue |
| 解析后守恒检查 | 数字字符守恒、文本覆盖率、表格还原度 | 超阈值 → 拒绝入库 |

核心原则：**宁可拒绝入库，也不能带病入库。** 解析错误不会让系统报错，
只会让答案悄悄变差，且事后无法定位到是哪份文档的哪次解析出了问题。

模块内不得出现数字字面量——所有阈值从 `QualityConfig` 读取。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from ..config import Config, QualityConfig
from ..schema import Block, ParsedDocument, QualityReport
from ..utils import digit_count
from .pdf_backend import get_pdf_backend

_TEXT_EXTS = {".md", ".markdown", ".txt"}
_TABLE_LINE = re.compile(r"^\s*\|")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


@dataclass
class ProbeResult:
    ok: bool = True
    encrypted: bool = False
    corrupted: bool = False
    is_scanned: bool = False
    chars_per_page: float = 0.0
    table_density: float = 0.0
    suggested_parser: str = ""


# ─────────────────────────── 关卡 0 · raw 基线 ───────────────────────────


def raw_baseline(path: Path, cfg: Config) -> str:
    """各格式的"最粗暴但完整"的基线文本（03 §4.3.0）。

    与解析器**不同源**才能查出漏项：DOCX/XLSX/PPTX/HTML 用的是另一套库。
    PDF 的 raw 与 PdfParser 同源（只是调用方式不同），这条检查的强度弱于其他格式，
    只能查出版面分析阶段丢掉的内容——这一点在文档里已写明。
    """
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            texts = get_pdf_backend(cfg).page_texts(path)
            return "\n".join(texts)
        if ext in (".docx", ".doc"):
            return _raw_docx(path)
        if ext in (".xlsx", ".xls"):
            return _raw_xlsx(path)
        if ext in (".pptx", ".ppt"):
            return _raw_pptx(path)
        if ext in (".html", ".htm"):
            return _raw_html(path)
        if ext in _TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:  # noqa: BLE001 - 基线取不到是预期情况，交给适用性判据处理
        return ""
    return ""


def _raw_docx(path: Path) -> str:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _raw_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts += [str(v) for v in row if v is not None]
    return "\n".join(parts)


def _raw_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _raw_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n")


def count_tables_detected(path: Path, cfg: Config) -> int:
    """独立于解析器的表格数量估计（表格还原度的分母）。"""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return get_pdf_backend(cfg).count_tables(path)
        if ext in (".docx", ".doc"):
            import docx

            return len(docx.Document(str(path)).tables)
        if ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True)
            return max(1, len(wb.worksheets))
        if ext in (".pptx", ".ppt"):
            from pptx import Presentation

            n = 0
            for slide in Presentation(str(path)).slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_table", False):
                        n += 1
            return n
        if ext in (".html", ".htm"):
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
            return len(soup.find_all("table"))
        if ext in _TEXT_EXTS:
            return len(_TABLE_LINE.findall(path.read_text(encoding="utf-8", errors="ignore"))) > 0 and 1 or 0
    except Exception:  # noqa: BLE001
        return 0
    return 0


# ─────────────────────────── 关卡 1 · 解析前探测 ───────────────────────────


def probe_before_parse(path: Path, cfg: Config) -> ProbeResult:
    path = Path(path)
    ext = path.suffix.lower()
    q = cfg.quality
    res = ProbeResult(suggested_parser=_suggest_parser(ext))

    if ext != ".pdf":
        try:
            text = raw_baseline(path, cfg)
        except Exception:  # noqa: BLE001
            text = ""
        res.chars_per_page = float(len(text))
        res.corrupted = not path.exists() or path.stat().st_size == 0
        res.ok = not res.corrupted
        return res

    backend = get_pdf_backend(cfg)
    try:
        pages = backend.page_texts(path)
    except Exception as e:  # noqa: BLE001
        res.ok = False
        res.corrupted = True
        res.suggested_parser = ""
        return res

    n_pages = max(1, len(pages))
    total = sum(len(p) for p in pages)
    res.chars_per_page = total / n_pages
    res.is_scanned = res.chars_per_page < q.scanned_chars_per_page
    res.encrypted = backend.is_encrypted(path)
    res.table_density = backend.table_density(path)
    res.corrupted = n_pages == 0 or (total == 0 and not res.is_scanned)
    res.ok = not (res.encrypted or res.corrupted)
    return res


def _suggest_parser(ext: str) -> str:
    return {
        ".pdf": "pymupdf4llm",
        ".docx": "markitdown",
        ".doc": "markitdown",
        ".xlsx": "markitdown",
        ".xls": "markitdown",
        ".pptx": "markitdown",
        ".ppt": "markitdown",
        ".html": "markitdown",
        ".htm": "markitdown",
        ".md": "native",
        ".markdown": "native",
        ".txt": "native",
    }.get(ext, "")


# ─────────────────────────── 关卡 2 · 结构校验 ───────────────────────────


def validate_structure(doc: ParsedDocument, cfg: QualityConfig) -> list[str]:
    issues: list[str] = []
    report = doc.quality

    # 表格行列一致
    ok_rows = tot_rows = 0
    tables = [b for b in doc.blocks if b.type == "table"]
    for b in tables:
        rows = _table_rows(b.text)
        if not rows:
            issues.append("table_unparsable")
            continue
        ncols = len(rows[0])
        for r in rows:
            tot_rows += 1
            if len(r) == ncols:
                ok_rows += 1
        if any(len(r) != ncols for r in rows):
            issues.append("table_row_col_mismatch")
    report.table_integrity = (ok_rows / tot_rows) if tot_rows else 1.0
    report.table_count_parsed = len(tables)

    # 标题层级连续
    levels = [b.level for b in doc.blocks if b.type == "heading" and b.level > 0]
    jumps = sum(1 for a, b in zip(levels, levels[1:]) if b - a > 1)
    report.heading_jumps = jumps
    if jumps > cfg.heading_jump_warn:
        issues.append("heading_level_jump")

    # 图文关联：孤立 caption
    for i, b in enumerate(doc.blocks):
        if b.type != "image_caption":
            continue
        prev = doc.blocks[i - 1] if i > 0 else None
        if prev is None or prev.type not in ("heading", "paragraph", "table"):
            issues.append("orphan_image_caption")

    report.issues.extend(issues)
    return issues


def _table_rows(text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in (text or "").splitlines():
        s = line.strip()
        if not _TABLE_LINE.match(s):
            continue
        if re.match(r"^\s*\|[\s:|-]+\|\s*$", s):  # 分隔行
            continue
        rows.append([c.strip() for c in s.strip("|").split("|")])
    return rows


def split_table_header(text: str) -> tuple[str, list[str]]:
    """返回 (表头行, 数据行列表)，供按行组切分时重复表头使用。"""
    rows = [line for line in (text or "").splitlines() if _TABLE_LINE.match(line.strip())]
    if not rows:
        return "", []
    sep_idx = None
    for i, line in enumerate(rows):
        if re.match(r"^\s*\|[\s:|-]+\|\s*$", line.strip()):
            sep_idx = i
            break
    if sep_idx is None:
        return rows[0], rows[1:]
    header = rows[: sep_idx + 1]
    return "\n".join(header), rows[sep_idx + 1 :]


# ─────────────────────────── 关卡 3 · 守恒检查 ───────────────────────────


def check_conservation(doc: ParsedDocument, raw: str, cfg: QualityConfig) -> list[str]:
    """守恒检查。**不适用时必须自我声明**，而不是算出一个假分数。"""
    report = doc.quality
    issues: list[str] = []
    parsed_text = "\n".join(b.text for b in doc.blocks)

    report.char_count = len(parsed_text)
    report.raw_char_count = len(raw or "")
    report.digit_count_parsed = digit_count(parsed_text)
    report.digit_count_raw = digit_count(raw or "")

    # md / txt 的 raw 与解析结果同源，守恒检查自动降级为不适用
    if Path(doc.source_path).suffix.lower() in _TEXT_EXTS:
        report.not_applicable.append("digit_loss")
        report.not_applicable.append("text_coverage")
        report.digit_loss = None
        report.text_coverage = None
    else:
        report.digit_loss = _digit_loss(report, cfg)
        report.text_coverage = _text_coverage(report, cfg)
        if report.digit_loss is None:
            report.not_applicable.append("digit_loss")
        elif report.digit_loss > cfg.digit_loss_warn:
            issues.append("digit_loss_warn")
        if report.text_coverage is None:
            report.not_applicable.append("text_coverage")
        elif report.text_coverage < cfg.text_coverage_warn:
            issues.append("text_coverage_warn")

    report.table_recall = _table_recall(report)
    if report.table_recall is None:
        report.not_applicable.append("table_recall")
    elif report.table_recall < cfg.table_recall_warn:
        issues.append("table_recall_warn")

    report.issues.extend(issues)
    return issues


def _digit_loss(report: QualityReport, cfg: QualityConfig) -> float | None:
    if report.digit_count_raw < cfg.digit_count_min_for_check:
        return None  # 原文数字量太少，比值无统计意义
    return 1 - report.digit_count_parsed / report.digit_count_raw


def _text_coverage(report: QualityReport, cfg: QualityConfig) -> float | None:
    if report.raw_char_count < cfg.raw_char_min_for_check:
        return None
    return report.char_count / report.raw_char_count


def _table_recall(report: QualityReport) -> float | None:
    if report.table_count_detected < 1:
        return None
    return report.table_count_parsed / report.table_count_detected


# ─────────────────────────── 汇总与处置 ───────────────────────────


def build_report(
    doc: ParsedDocument,
    raw: str,
    cfg: QualityConfig,
    ocr_available: bool,
    probe: ProbeResult | None = None,
) -> QualityReport:
    report = doc.quality
    report.parser = doc.parser
    report.page_count = doc.page_count
    if probe is not None:
        report.is_scanned = probe.is_scanned
    if not report.char_count:
        report.char_count = sum(len(b.text) for b in doc.blocks)
    _verdict(report, cfg, ocr_available)
    return report


def _verdict(report: QualityReport, cfg: QualityConfig, ocr_available: bool) -> str:
    # 1. 扫描件但没有可用 OCR 通道 → 拒绝（判定发生在路由降级之后）
    if report.is_scanned and not ocr_available:
        report.issues.append("scanned_without_ocr")
        report.verdict = "rejected"
        return report.verdict
    # 2. 扫描件走了 OCR → 守恒检查不适用，改看绝对量下限
    if report.is_scanned and ocr_available and report.char_count < cfg.ocr_min_chars:
        report.issues.append("ocr_output_too_small")
        report.verdict = "rejected"
        return report.verdict
    # 3. 严重项（仅当检查适用时才参与判定）
    #    拒绝原因必须显式写进 issues：下游 `pipeline` 用 issues[-1] 当拒答理由报给运维，
    #    空着的话报出去的是最后一个无关的轻微项（标题跳级、表格行列对不上…），
    #    真正的拒因（数字丢失率超阈值）反而消失——排查时会被带到完全错误的方向。
    if report.digit_loss is not None and report.digit_loss > cfg.digit_loss_reject:
        report.issues.append(
            f"digit_loss_reject:{round(report.digit_loss, 4)}>{cfg.digit_loss_reject}"
        )
        report.verdict = "rejected"
        return report.verdict
    if report.text_coverage is not None and report.text_coverage < cfg.text_coverage_reject:
        report.issues.append(
            f"text_coverage_reject:{round(report.text_coverage, 4)}<{cfg.text_coverage_reject}"
        )
        report.verdict = "rejected"
        return report.verdict
    # 4. 一般项
    report.verdict = "warned" if report.issues else "ok"
    return report.verdict


@dataclass
class DocumentQuality:
    """一次文档质量检查的完整产物（`inspect` 命令直接展示它）。"""

    doc_id: str = ""
    source: str = ""
    report: QualityReport = field(default_factory=QualityReport)
    raw_chars: int = 0


def markdown_to_blocks(md: str) -> list[Block]:
    """Markdown → 统一中间表示（native 解析器与 markitdown 输出共用）。"""
    blocks: list[Block] = []
    path: list[str] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush(kind: str = "paragraph") -> None:
        if not buf:
            return
        text = "\n".join(buf).strip()
        buf.clear()
        if text:
            blocks.append(Block(type=kind, text=text, path=list(path)))

    in_code = False
    for line in (md or "").splitlines():
        if line.strip().startswith("```"):
            if in_code:
                flush("code")
                in_code = False
            else:
                flush()
                in_code = True
            continue
        if in_code:
            buf.append(line)
            continue

        m = _HEADING.match(line.strip())
        if m:
            flush()
            level = len(m.group(1))
            title = m.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            path = [t for _, t in stack]
            blocks.append(Block(type="heading", text=title, level=level, path=list(path)))
            continue

        if _TABLE_LINE.match(line):
            if blocks and blocks[-1].type == "table":
                blocks[-1].text += "\n" + line.rstrip()
            else:
                flush()
                blocks.append(
                    Block(type="table", text=line.rstrip(), path=list(path), is_atomic=True)
                )
            continue

        s = line.strip()
        if not s:
            flush()
            continue
        if re.match(r"^!\[.*\]\(.*\)\s*$", s):
            flush()
            blocks.append(Block(type="image_caption", text=s, path=list(path)))
            continue
        if re.match(r"^([-*+]|\d+[.)])\s+", s):
            if blocks and blocks[-1].type == "list":
                blocks[-1].text += "\n" + s
            else:
                flush()
                blocks.append(Block(type="list", text=s, path=list(path)))
            continue

        buf.append(line)

    flush("code" if in_code else "paragraph")
    return blocks


__all__ = [
    "ProbeResult",
    "probe_before_parse",
    "validate_structure",
    "raw_baseline",
    "check_conservation",
    "build_report",
    "count_tables_detected",
    "markdown_to_blocks",
    "split_table_header",
]
