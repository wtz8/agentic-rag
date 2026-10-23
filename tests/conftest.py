"""测试公共 fixture（L1–L3 必须能在无 API Key 的机器上全绿）。"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
TESTS = ROOT / "tests"
for p in (SRC, TESTS):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))


@pytest.fixture
def cfg():
    """Mock 模式配置（无 API Key）。"""
    from agentic_rag.config import get_config

    return get_config(env={"DASHSCOPE_API_KEY": ""})


def pytest_configure(config: pytest.Config) -> None:
    """让 `--basetemp=<子目录>` 在干净克隆上也能用。

    pytest 只对 `--basetemp` 的**末级目录**做 `mkdir(parents=False)`，父目录不存在时
    直接抛 FileNotFoundError（报错发生在任何用例之前，看起来像"整包测试坏了"）。
    本仓库原先能跑，只是因为工作区里恰好留着 `.pytest_tmp/`；新克隆的仓库没有它，
    README 与 CI 用的 `.pytest_tmp/ci` 会当场失败。这里补出父目录，路径由调用方决定。
    """
    basetemp = getattr(config.option, "basetemp", None)
    if not basetemp:
        return
    Path(basetemp).expanduser().resolve().parent.mkdir(parents=True, exist_ok=True)


@pytest.fixture
def sample_chunks():
    """5 篇微型语料构造的 chunk，用于检索侧测试。"""
    from agentic_rag.schema import Chunk

    texts = [
        ("华东区带宽采购单价", "华东区 2024 年带宽采购单价为 3.2 万元/Gbps。"),
        ("华北区带宽采购单价", "华北区 2024 年带宽采购单价为 2.8 万元/Gbps。"),
        ("设备型号 XR-2400", "设备型号 XR-2400 支持双电源冗余，最大吞吐 400Gbps。"),
        ("错误码 E0107", "错误码 E0107 表示光模块收光功率低于阈值。"),
        ("研发投入", "2024 年研发投入 1.2 亿元，占营收 8.5%。"),
    ]
    out = []
    for i, (title, body) in enumerate(texts):
        out.append(
            Chunk(
                chunk_id=f"doc{i:02d}:c{i:04d}",
                doc_id=f"doc{i:02d}",
                parent_id=None,
                text=f"[t{i}.txt > {title}]\n{body}",
                raw_text=body,
                title_path=[title],
                block_type="paragraph",
                page=None,
                source=f"t{i}.txt",
                tokens=len(body),
            )
        )
    return out


# ── 语料 fixture：覆盖 8 种格式 ────────────────────────────────────────


def _write_md(p: Path) -> Path:
    p.write_text(
        "# 一级标题\n\n正文段落一，包含数字 12345。\n\n## 二级标题\n\n"
        "| 区域 | 单价 |\n| --- | --- |\n| 华东 | 3.2 |\n| 华北 | 2.8 |\n\n"
        "### 三级标题\n\n正文段落二 67890。\n",
        encoding="utf-8",
    )
    return p


def _write_txt(p: Path) -> Path:
    p.write_text("纯文本语料。\n包含数字 100 与 200。\n" * 30, encoding="utf-8")
    return p


def _write_html(p: Path) -> Path:
    p.write_text(
        "<html><body><h1>标题一</h1><p>段落文字 1234。</p>"
        "<table><tr><th>项</th><th>值</th></tr><tr><td>带宽</td><td>3.2</td></tr></table>"
        "<h2>标题二</h2><ul><li>要点一</li><li>要点二</li></ul></body></html>",
        encoding="utf-8",
    )
    return p


def _write_docx(p: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("一级标题", level=1)
    d.add_paragraph("正文段落，含数字 12345。")
    d.add_heading("二级标题", level=2)
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "区域"
    t.cell(0, 1).text = "单价"
    t.cell(1, 0).text = "华东"
    t.cell(1, 1).text = "3.2"
    t.cell(2, 0).text = "华北"
    t.cell(2, 1).text = "2.8"
    d.save(str(p))
    return p


def _write_xlsx(p: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["区域", "单价", "数量"])
    ws.append(["华东", 3.2, 100])
    ws.append(["华北", 2.8, 200])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["年份", "营收"])
    ws2.append([2024, 120000])
    wb.save(str(p))
    return p


def _write_pptx(p: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "部署架构"
    s.placeholders[1].text = "要点一：双机热备\n要点二：同城双活"
    prs.save(str(p))
    return p


def _write_pdf(p: Path, pages: int = 2, text: str | None = None) -> Path:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 100), text or f"第{i + 1}页正文内容，含数字 12345 与 67890。" * 3)
    doc.save(str(p))
    doc.close()
    return p


def _write_scanned_pdf(p: Path) -> Path:
    """无文本层的 PDF：只画一个矩形，不插文字。"""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.draw_rect(fitz.Rect(72, 100, 500, 300), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    doc.save(str(p))
    doc.close()
    return p


@pytest.fixture
def corpus_dir(tmp_path):
    d = tmp_path / "raw"
    d.mkdir()
    _write_md(d / "a.md")
    _write_txt(d / "b.txt")
    _write_html(d / "c.html")
    _write_docx(d / "d.docx")
    _write_xlsx(d / "e.xlsx")
    _write_pptx(d / "f.pptx")
    _write_pdf(d / "g.pdf")
    return d


@pytest.fixture
def make_file(tmp_path):
    def _make(kind: str, name: str | None = None, **kw):
        return {
            "md": _write_md,
            "txt": _write_txt,
            "html": _write_html,
            "docx": _write_docx,
            "xlsx": _write_xlsx,
            "pptx": _write_pptx,
            "pdf": _write_pdf,
            "scanned_pdf": _write_scanned_pdf,
        }[kind](tmp_path / (name or f"t.{kind if kind != 'scanned_pdf' else 'pdf'}"), **kw)

    return _make


# ── 集成向 fixture：把落盘路径全部重定向到 tmp_path ──────────────────────


@pytest.fixture
def icfg(cfg, tmp_path):
    """集成测试的配置：一切落盘路径都在 tmp，绝不污染仓库。"""
    cfg.index.path = str(tmp_path / "index")
    cfg.index.backend = "memory"
    cfg.observability.trace_dir = str(tmp_path / "traces")
    cfg.cache.dir = str(tmp_path / "cache")
    cfg.ingestion.workers = 2
    return cfg


@pytest.fixture
def store_with_corpus(icfg, corpus_dir):
    """真实摄取 7 种格式 → 真实索引。"""
    from agentic_rag.index.store import IndexStore
    from agentic_rag.ingestion.pipeline import load_documents

    chunks, _ = load_documents(corpus_dir, icfg, move_rejected=False)
    st = IndexStore(icfg)
    st.build(chunks)
    return st


@pytest.fixture
def tiny_store(icfg):
    """只有 2 个 chunk 的索引：把"候选数永远不够"变成确定性前提。"""
    from agentic_rag.index.store import IndexStore
    from agentic_rag.schema import Chunk

    texts = [
        "华东区 2024 年带宽采购单价为 3.2 万元/Gbps。",
        "华北区 2024 年带宽采购单价为 2.8 万元/Gbps。",
    ]
    chunks = [
        Chunk(
            chunk_id=f"d{i}:c0000",
            doc_id=f"d{i}",
            parent_id=None,
            text=t,
            raw_text=t,
            title_path=["带宽采购"],
            source=f"tiny{i}.md",
            tokens=20,
        )
        for i, t in enumerate(texts)
    ]
    st = IndexStore(icfg)
    st.build(chunks)
    return st


@pytest.fixture
def pipeline_factory(icfg):
    """组装 Pipeline，允许替换模型层（集成测试要断言的是编排路径，不是模型行为）。"""
    from agentic_rag.bootstrap import build_deps
    from agentic_rag.pipeline import Pipeline

    def _make(store, llm_light=None, llm_main=None):
        deps = build_deps(icfg, store)
        if llm_light is not None:
            deps.llm_light = llm_light
        if llm_main is not None:
            deps.llm_main = llm_main
        return Pipeline(icfg, deps, store)

    return _make
